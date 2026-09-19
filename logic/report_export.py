"""Generic table exporters for reports: Excel (openpyxl) and PDF (reportlab)."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


def export_table_to_excel(title: str, headers: list[str], rows: list[list], file_path: str) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31] or "Report"

    ws.append(headers)
    header_fill = PatternFill(start_color="1565C0", end_color="1565C0", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font

    for row in rows:
        ws.append(row)

    for column_cells in ws.columns:
        length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
        ws.column_dimensions[column_cells[0].column_letter].width = min(max(length + 2, 10), 40)

    path = Path(file_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def export_table_to_pdf(title: str, headers: list[str], rows: list[list], file_path: str) -> Path:
    path = Path(file_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(path), pagesize=landscape(A4), topMargin=14 * mm, bottomMargin=14 * mm)
    elements = [Paragraph(f"<b>{title}</b>", styles["Title"]), Spacer(1, 6 * mm)]

    data = [headers] + [[str(v) for v in row] for row in rows]
    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1565C0")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F7FA")]),
            ]
        )
    )
    elements.append(table)
    doc.build(elements)
    return path
